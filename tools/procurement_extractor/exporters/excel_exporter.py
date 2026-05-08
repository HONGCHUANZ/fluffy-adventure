from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _safe_sheet_title(title: str, used_titles: set[str]) -> str:
    sanitized = title.strip() or "Sheet"
    sanitized = sanitized.translate(str.maketrans({
        "/": "_",
        "\\": "_",
        "?": "_",
        "*": "_",
        "[": "(",
        "]": ")",
        ":": "-",
    }))
    sanitized = sanitized[:31] or "Sheet"
    candidate = sanitized
    counter = 2
    while candidate in used_titles:
        suffix = f"_{counter}"
        candidate = f"{sanitized[:31 - len(suffix)]}{suffix}"
        counter += 1
    used_titles.add(candidate)
    return candidate


def _populate_sheet(worksheet, records: list[dict[str, str]], columns: list[tuple[str, str]]) -> None:
    headers = [label for _, label in columns]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for record in records:
        worksheet.append([record.get(name, "") for name, _ in columns])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for index, header in enumerate(headers, start=1):
        max_length = len(header)
        for row in worksheet.iter_rows(min_row=2, min_col=index, max_col=index):
            value = "" if row[0].value is None else str(row[0].value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 80)


def export_records(site_exports: list[dict[str, object]], output_path: str | Path, *, target_date: str) -> Path:
    workbook = Workbook()
    initial_sheet = workbook.active
    workbook.remove(initial_sheet)

    used_titles: set[str] = set()
    total_records = 0

    for site_export in site_exports:
        records = site_export["records"]
        columns = site_export["columns"]
        sheet_name = site_export["sheet_name"]
        worksheet = workbook.create_sheet(_safe_sheet_title(str(sheet_name), used_titles))
        _populate_sheet(worksheet, records, columns)
        total_records += len(records)

    summary = workbook.create_sheet("run_summary")
    summary.append(["目标日期", target_date])
    summary.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    summary.append(["站点数量", len(site_exports)])
    summary.append(["公告总数", total_records])
    summary.append([])
    summary.append(["站点标识", "站点名称", "工作表", "公告数量"])

    for site_export in site_exports:
        summary.append(
            [
                str(site_export["site_key"]),
                str(site_export["site_name"]),
                str(site_export["sheet_name"]),
                len(site_export["records"]),
            ]
        )

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
